"""Reading a .docx into segments, with enough structure to write back into it.

Two things make this more than ``document.paragraphs``:

**Document order.** ``python-docx`` exposes paragraphs and tables as separate
collections, so neither preserves the order text appears on the page. This reader
walks the body's XML children instead, which keeps paragraphs, tables and nested
tables in reading order. That matters for the audit log (locations are cited in
document order) and for section attribution.

**Run-level offset maps.** A redaction has to be written back into the exact runs
it came from or all character formatting is lost. Each segment therefore carries
a list of :class:`RunRef` entries recording which slice of the segment's text each
``w:r`` element owns. Word splits text across runs unpredictably -- a single
company name is routinely three runs because of a spell-check boundary or a
tracked-change remnant -- so this mapping, not the text, is the reader's real
output.

Table cells additionally carry their column header and row label; see
:class:`piiredact.types.Segment` for why that is worth the trouble.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Iterator

import docx
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from ..types import Segment

#: Paragraphs are joined by this character inside a multi-paragraph segment.
PARAGRAPH_SEPARATOR = "\n"


@dataclass(slots=True)
class RunRef:
    """One ``w:r`` run and the slice of segment text it owns."""

    run: Run
    start: int
    end: int
    #: Index of the paragraph this run belongs to within the segment.
    paragraph_index: int

    def __len__(self) -> int:
        return self.end - self.start


@dataclass(slots=True)
class SegmentHandle:
    """A segment plus everything needed to rewrite it in place."""

    segment: Segment
    runs: list[RunRef] = field(default_factory=list)
    paragraphs: list[Paragraph] = field(default_factory=list)


@dataclass(slots=True)
class ReadResult:
    """What :func:`read_document` produces."""

    document: DocxDocument
    handles: list[SegmentHandle]

    @property
    def segments(self) -> list[Segment]:
        return [handle.segment for handle in self.handles]

    def by_key(self) -> dict[str, SegmentHandle]:
        return {handle.segment.key: handle for handle in self.handles}

    def stats(self) -> dict[str, int]:
        paragraph_segments = sum(1 for h in self.handles if h.segment.kind == "paragraph")
        cell_segments = sum(1 for h in self.handles if h.segment.is_table_cell)
        return {
            "segments": len(self.handles),
            "paragraph_segments": paragraph_segments,
            "table_cell_segments": cell_segments,
            "runs": sum(len(h.runs) for h in self.handles),
            "characters": sum(len(h.segment.text) for h in self.handles),
        }


#: Wrappers whose ``w:r`` descendants are *not* displayed text and must be left
#: alone: deleted tracked-changes content, and the original of a moved range.
_SKIP_ANCESTORS = frozenset({qn("w:del"), qn("w:moveFrom")})


def _paragraph_runs(paragraph: Paragraph) -> list[Run]:
    """Every displayed run of a paragraph, in document order.

    ``Paragraph.runs`` returns only the paragraph's *direct* ``w:r`` children, so
    it silently skips runs nested inside ``w:hyperlink``, ``w:sdt`` (content
    controls), ``w:ins`` (tracked insertions) and ``w:smartTag``. Those runs are
    displayed text, and missing them means the detectors never see the characters
    and the writer never rewrites them -- which is exactly how one shareholder
    name survived a full run of this tool before this function was made
    exhaustive. Walking all ``w:r`` descendants and skipping only the wrappers
    that hold undisplayed text is both simpler and correct.
    """
    element = paragraph._p
    runs: list[Run] = []
    for r in element.iter(qn("w:r")):
        ancestor = r.getparent()
        skip = False
        while ancestor is not None and ancestor is not element:
            if ancestor.tag in _SKIP_ANCESTORS:
                skip = True
                break
            ancestor = ancestor.getparent()
        if not skip:
            runs.append(Run(r, paragraph))
    return runs


def _build_handle(
    key: str,
    paragraphs: list[Paragraph],
    location: str,
    kind: str,
    *,
    section: str | None = None,
    style: str | None = None,
    table_index: int | None = None,
    row_index: int | None = None,
    col_index: int | None = None,
    column_header: str | None = None,
    row_header: str | None = None,
) -> SegmentHandle | None:
    """Assemble one segment and its run map, or None when there is no text."""
    parts: list[str] = []
    runs: list[RunRef] = []
    cursor = 0
    for paragraph_index, paragraph in enumerate(paragraphs):
        if paragraph_index:
            parts.append(PARAGRAPH_SEPARATOR)
            cursor += len(PARAGRAPH_SEPARATOR)
        for run in _paragraph_runs(paragraph):
            run_text = run.text
            if not run_text:
                continue
            runs.append(RunRef(run, cursor, cursor + len(run_text), paragraph_index))
            parts.append(run_text)
            cursor += len(run_text)
    text = "".join(parts)
    if not text.strip():
        return None
    segment = Segment(
        key=key,
        text=text,
        location=location,
        kind=kind,
        section=section,
        style=style,
        table_index=table_index,
        row_index=row_index,
        col_index=col_index,
        column_header=column_header,
        row_header=row_header,
    )
    return SegmentHandle(segment=segment, runs=runs, paragraphs=list(paragraphs))


def _cell_text(cell: _Cell) -> str:
    return " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip()).strip()


def _iter_block_items(parent: object) -> Iterator[Paragraph | Table]:
    """Yield paragraphs and tables of ``parent`` in document order."""
    if isinstance(parent, _Cell):
        element = parent._tc
    else:
        element = parent.element.body  # type: ignore[union-attr]
    for child in element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)  # type: ignore[arg-type]
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)  # type: ignore[arg-type]


class _Reader:
    """Stateful walk of the document body."""

    def __init__(self, document: DocxDocument) -> None:
        self.document = document
        self.handles: list[SegmentHandle] = []
        self.section: str | None = None
        self._paragraph_counter = 0
        self._table_counter = 0

    def run(self) -> ReadResult:
        self._walk(self.document, path="body")
        return ReadResult(document=self.document, handles=self.handles)

    def _walk(self, parent: object, path: str) -> None:
        for item in _iter_block_items(parent):
            if isinstance(item, Paragraph):
                self._handle_paragraph(item, path)
            else:
                self._handle_table(item, path)

    # -- paragraphs ----------------------------------------------------------

    def _handle_paragraph(self, paragraph: Paragraph, path: str) -> None:
        style_name = (paragraph.style.name if paragraph.style is not None else "") or ""
        text = paragraph.text.strip()
        if text and style_name.lower().startswith("heading"):
            self.section = text
        index = self._paragraph_counter
        self._paragraph_counter += 1
        handle = _build_handle(
            key=f"{path}/p{index}",
            paragraphs=[paragraph],
            location=f"paragraph {index}",
            kind="paragraph",
            section=self.section,
            style=style_name or None,
        )
        if handle is not None:
            self.handles.append(handle)

    # -- tables --------------------------------------------------------------

    def _handle_table(self, table: Table, path: str) -> None:
        table_index = self._table_counter
        self._table_counter += 1
        table_path = f"{path}/tbl{table_index}"

        rows = list(table.rows)
        header_cells = list(rows[0].cells) if rows else []
        headers = [_cell_text(cell) for cell in header_cells]

        # Merged cells repeat the same ``w:tc`` across grid positions, so each
        # underlying cell must be processed exactly once. The key is the element's
        # XML path rather than ``id(cell._tc)``: lxml hands out *transient*
        # proxies, so a freshly collected proxy can reuse an earlier object's
        # ``id`` and make an unrelated cell look already-seen. That happened here
        # -- it silently dropped one shareholder-name cell from the whole run.
        seen_cells: set[str] = set()
        for row_index, row in enumerate(rows):
            cells = list(row.cells)
            row_header = _cell_text(cells[0]) if cells else ""
            for col_index, cell in enumerate(cells):
                cell_key = cell._tc.getroottree().getpath(cell._tc)
                if cell_key in seen_cells:
                    continue
                seen_cells.add(cell_key)

                column_header = headers[col_index] if col_index < len(headers) else None
                if row_index == 0:
                    column_header = None  # the header cell has no header of its own

                handle = _build_handle(
                    key=f"{table_path}/r{row_index}c{col_index}",
                    paragraphs=list(cell.paragraphs),
                    location=f"table {table_index}, row {row_index}, col {col_index}",
                    kind="table-cell",
                    section=self.section,
                    table_index=table_index,
                    row_index=row_index,
                    col_index=col_index,
                    column_header=column_header,
                    row_header=row_header if col_index else None,
                )
                if handle is not None:
                    self.handles.append(handle)

                # Nested tables inside this cell.
                for child in _iter_block_items(cell):
                    if isinstance(child, Table):
                        self._handle_table(child, f"{table_path}/r{row_index}c{col_index}")


def read_document(path: str) -> ReadResult:
    """Open ``path`` and extract every text segment with its run map."""
    document = docx.Document(path)
    return _Reader(document).run()


__all__ = ["ReadResult", "RunRef", "SegmentHandle", "read_document", "PARAGRAPH_SEPARATOR"]
